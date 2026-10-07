import re
from datetime import datetime
import openpyxl
import requests

EXCEL_FILE = "Job applications list.xlsx"
TARGET_SHEET_NAME = "Automated Jobs 2026-2027"

# Internship Data Sources
SOURCES = [
    {
        "name": "Canada Tech Internships",
        "urls": [
            "https://raw.githubusercontent.com/negarprh/Canadian-Tech-Internships-2027/main/README.md",
        ],
    },
    {
        "name": "SimplifyJobs Summer 2027",
        "urls": [
            "https://raw.githubusercontent.com/SimplifyJobs/Summer2027-Internships/dev/README.md",
            "https://raw.githubusercontent.com/SimplifyJobs/Summer2027-Internships/main/README.md",
        ],
    },
]


def normalize_github_url(raw_url):
    """Normalizes GitHub web/raw URLs into a valid raw.githubusercontent.com endpoint."""
    url = raw_url.strip()
    if "github.com" in url and "raw.githubusercontent.com" not in url:
        url = url.replace("github.com", "raw.githubusercontent.com")
        url = url.replace("/blob/", "/")
        url = url.replace("/raw/", "/")

    # Strip refs/heads/ from raw.githubusercontent.com paths
    url = re.sub(
        r"(raw\.githubusercontent\.com/[^/]+/[^/]+/)refs/heads/", r"\1", url
    )
    return url


def extract_application_url(text):
    """Extracts application URL while filtering out shield/badge image URLs."""
    if not text:
        return ""
    urls = re.findall(r'href=["\'](.*?)["\']', text) + re.findall(
        r"\((https?://.*?)\)", text
    )
    valid_urls = [
        u.strip()
        for u in urls
        if not any(
            ignore in u.lower()
            for ignore in [
                "shields.io",
                "badge",
                ".png",
                ".svg",
                ".jpg",
                ".jpeg",
                ".gif",
            ]
        )
    ]
    return valid_urls[0] if valid_urls else ""


def clean_text(text):
    """Strips HTML tags and Markdown formatting while keeping clean text."""
    if not text:
        return ""
    text = re.sub(r"<[^>]+>", "", text)
    text = re.sub(r"\[(.*?)\]\(.*?\)", r"\1", text)
    text = re.sub(r"[*_]+", "", text)
    return text.strip()


def parse_markdown_table(markdown_text, source_name):
    """Parses markdown tables using dynamic header mapping."""
    jobs = []
    lines = markdown_text.splitlines()

    col_map = {}
    last_company = ""

    for line in lines:
        line_str = line.strip()
        if not line_str.startswith("|") or "|" not in line_str:
            continue

        parts = [p.strip() for p in line_str.split("|")][1:-1]
        if not parts:
            continue

        header_line_lower = [p.lower() for p in parts]

        # Detect Header Row to dynamically map columns
        if any(
            h in header_line_lower
            for h in ["company", "organization", "employer"]
        ):
            col_map = {}
            for idx, h in enumerate(header_line_lower):
                if any(
                    k in h for k in ["company", "organization", "employer"]
                ):
                    col_map["company"] = idx
                elif any(k in h for k in ["role", "title", "position"]):
                    col_map["title"] = idx
                elif any(k in h for k in ["location", "city"]):
                    col_map["location"] = idx
                elif any(k in h for k in ["app", "link", "apply"]):
                    col_map["application"] = idx
                elif any(k in h for k in ["date", "posted", "age"]):
                    col_map["date_posted"] = idx
            continue

        # Skip separator row (|---|---|---|)
        if all(re.match(r"^[\s\-:\*]+$", p) for p in parts if p):
            continue

        # Default fallback indices if header was omitted
        c_idx = col_map.get("company", 0)
        t_idx = col_map.get("title", 1)
        l_idx = col_map.get("location", 2)
        a_idx = col_map.get("application", 3)
        d_idx = col_map.get("date_posted", 4)

        raw_company = parts[c_idx] if len(parts) > c_idx else ""
        raw_title = parts[t_idx] if len(parts) > t_idx else ""
        raw_location = parts[l_idx] if len(parts) > l_idx else ""
        raw_app = parts[a_idx] if len(parts) > a_idx else ""
        raw_date = parts[d_idx] if len(parts) > d_idx else ""

        # Flag extraction for Notes field
        flags = []
        if "🇺🇸" in line_str:
            flags.append("🇺🇸 US Citizen Required")
        if "🛂" in line_str:
            flags.append("🛂 Visa Sponsorship Offered")
        if "🚫" in line_str:
            flags.append("🚫 No Sponsorship Offered")
        if "🎓" in line_str:
            flags.append("🎓 Advanced Degree Required")
        if "🔒" in line_str or "Closed" in raw_app or "Closed" in line_str:
            flags.append("🔒 Application Closed")

        is_closed = (
            "🔒" in line_str or "Closed" in raw_app or "Closed" in line_str
        )

        # Handle sub-roles (↳) and company memory
        clean_comp = clean_text(raw_company)
        clean_comp_alpha = re.sub(r"[^\w\s\-\.&,\(\)]", "", clean_comp).strip()

        if (
            "↳" in raw_company
            or "↳" in raw_title
            or not clean_comp_alpha
            or clean_comp_alpha.lower() in ["same", "ditto"]
        ):
            company = last_company
        else:
            company = clean_comp_alpha
            last_company = clean_comp_alpha

        clean_job_title = clean_text(raw_title)
        clean_job_title = re.sub(r"↳", "", clean_job_title)
        clean_job_title = re.sub(
            r"[^\w\s\-\.&,\(\)]", "", clean_job_title
        ).strip()

        # Clean location and date fields
        loc_text = raw_location.replace("<br>", " | ").replace("<br/>", " | ")
        clean_loc = clean_text(loc_text)
        clean_date = clean_text(raw_date)

        # Extract destination application link
        job_url = (
            extract_application_url(raw_app)
            or extract_application_url(raw_title)
            or extract_application_url(raw_company)
        )

        if company and clean_job_title:
            jobs.append({
                "company": company,
                "title": clean_job_title,
                "location": clean_loc or "Remote / Flexible",
                "url": job_url,
                "date_posted": clean_date or "N/A",
                "flags": " | ".join(flags) if flags else "None",
                "is_closed": is_closed,
                "source": source_name,
            })

    return jobs


def initialize_automated_sheet(wb):
    if TARGET_SHEET_NAME in wb.sheetnames:
        return wb[TARGET_SHEET_NAME]

    ws = wb.create_sheet(title=TARGET_SHEET_NAME)
    headers = [
        "Status",
        "Company",
        "Job Title",
        "Semester",
        "Location",
        "Date posted",
        "Date written in Excel",
        "Date applied",
        "Closing date",
        "Job post link",
        "Interest in the role",
        "Salary",
        "Notes",
    ]
    ws.append(headers)
    print(f"Created sheet '{TARGET_SHEET_NAME}' with headers.")
    return ws


def fetch_raw_markdown(urls):
    for raw_url in urls:
        normalized_url = normalize_github_url(raw_url)
        try:
            res = requests.get(normalized_url, timeout=10)
            if res.status_code == 200 and len(res.text) > 100:
                return res.text
        except Exception:
            continue
    return None


def sync_all_internships():
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = initialize_automated_sheet(wb)

    existing_links = set()
    existing_keys = set()
    row_index_map = {}

    for r in range(2, ws.max_row + 1):
        company = ws.cell(row=r, column=2).value
        title = ws.cell(row=r, column=3).value
        link = ws.cell(row=r, column=10).value

        if link:
            existing_links.add(str(link).strip().lower())
        if company and title:
            key = (str(company).strip().lower(), str(title).strip().lower())
            existing_keys.add(key)
            row_index_map[key] = r

    now = datetime.now()
    today_formatted = f"{now.month}/{now.day}/{now.year}"  # M/D/YYYY

    total_added = 0
    total_updated_closed = 0

    for source in SOURCES:
        print(f"Fetching {source['name']}...")
        md_text = fetch_raw_markdown(source["urls"])

        if not md_text:
            print(f"   ⚠️ Could not fetch valid Markdown for {source['name']}.")
            continue

        parsed_jobs = parse_markdown_table(md_text, source["name"])
        print(f"   ↳ Parsed {len(parsed_jobs)} postings from {source['name']}.")

        for job in parsed_jobs:
            key = (job["company"].lower(), job["title"].lower())
            link_lower = job["url"].lower()

            if job["is_closed"]:
                if key in row_index_map:
                    r = row_index_map[key]
                    if ws.cell(row=r, column=1).value == "Not applied":
                        ws.cell(row=r, column=1, value="Unavailable")
                        current_notes = ws.cell(row=r, column=13).value or ""
                        ws.cell(
                            row=r,
                            column=13,
                            value=f"{current_notes} | Marked unavailable on {today_formatted}",
                        )
                        total_updated_closed += 1
                continue

            if (link_lower and link_lower in existing_links) or (
                key in existing_keys
            ):
                continue

            notes_field = f"Source: {job['source']}"
            if job["flags"] != "None":
                notes_field += f" | Flags: {job['flags']}"

            new_row = [
                "Not applied",  # Col 1: Status
                job["company"],  # Col 2: Company
                job["title"],  # Col 3: Job Title
                "SUMMER 2027",  # Col 4: Semester
                job["location"],  # Col 5: Location
                job["date_posted"],  # Col 6: Date posted
                today_formatted,  # Col 7: Date written in Excel (M/D/YYYY)
                None,  # Col 8: Date applied
                None,  # Col 9: Closing date
                job["url"],  # Col 10: Clean Job post link
                "Medium",  # Col 11: Interest in the role
                None,  # Col 12: Salary
                notes_field,  # Col 13: Notes (Source + Flags)
            ]

            ws.append(new_row)
            existing_keys.add(key)
            if job["url"]:
                existing_links.add(link_lower)
            total_added += 1

    wb.save(EXCEL_FILE)
    print(f"\n✅ Sync Complete for sheet '{TARGET_SHEET_NAME}':")
    print(f"   • {total_added} new postings added.")
    print(
        f"   • {total_updated_closed} postings updated to 'Unavailable' status."
    )


if __name__ == "__main__":
    sync_all_internships()